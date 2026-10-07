"""
Backend REST API Server for DOI Monitoring Dashboard (MNJ & KX Principal)
Provides REST API endpoints for MNJ Stock, KX Principal Stock, Combined DOI Metrics, Multi-Select Filters, and Exports.
"""

import sys
import os
import json
import csv
import io
import mimetypes
import threading
from urllib.parse import parse_qs, urlparse
from http.server import HTTPServer, ThreadingHTTPServer, BaseHTTPRequestHandler

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from etl import DataEngine, parse_multi_param
from ppt_generator import generate_doi_ppt

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")

data_engine = DataEngine(base_dir=BASE_DIR)

class DOIRequestHandler(BaseHTTPRequestHandler):
    def _set_headers(self, status=200, content_type="application/json"):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, Authorization")
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.end_headers()

    def serve_file(self, full_path: str):
        full_path = os.path.normpath(full_path)
        if os.path.isfile(full_path):
            content_type, _ = mimetypes.guess_type(full_path)
            if not content_type:
                if full_path.endswith(".js"):
                    content_type = "application/javascript"
                elif full_path.endswith(".css"):
                    content_type = "text/css"
                elif full_path.endswith(".html"):
                    content_type = "text/html"
                else:
                    content_type = "application/octet-stream"

            if content_type.startswith("text/") or content_type == "application/javascript":
                content_type += "; charset=utf-8"

            self._set_headers(200, content_type=content_type)
            with open(full_path, "rb") as f:
                self.wfile.write(f.read())
        else:
            self._set_headers(404)
            self.wfile.write(json.dumps({"error": f"File not found: {full_path}"}).encode("utf-8"))

    def do_OPTIONS(self):
        self._set_headers(200)

    def do_HEAD(self):
        self._set_headers(200)

    def do_GET(self):
        parsed_url = urlparse(self.path)
        path = parsed_url.path
        query_params = parse_qs(parsed_url.query)

        def get_param(name: str, default: str) -> str:
            return query_params.get(name, [default])[0]

        try:
            if path == "/health":
                self._set_headers(200)
                res = {"status": "online", "version": "2.0.0", "app": "DOI MNJ & KX Monitoring API"}
                self.wfile.write(json.dumps(res).encode("utf-8"))

            elif path == "/api/v1/metadata":
                self.handle_metadata()

            elif path == "/api/v1/summary":
                self.handle_summary(get_param)

            elif path == "/api/v1/gb-summary":
                self.handle_gb_summary(get_param)

            elif path == "/api/v1/category-gb-summary":
                self.handle_category_gb_summary(get_param)

            elif path == "/api/v1/doi-trend":
                self.handle_doi_trend(get_param)

            elif path == "/api/v1/doi-data":
                self.handle_doi_data(get_param)

            elif path == "/api/v1/export":
                self.handle_export(get_param)

            elif path == "/api/v1/export-excel":
                self.handle_export_excel(get_param)

            elif path == "/api/v1/export-ppt":
                self.handle_export_ppt(get_param)

            else:
                rel_path = path.lstrip("/")
                if not rel_path or rel_path == "index.html":
                    target_file = os.path.join(FRONTEND_DIR, "index.html")
                else:
                    target_file = os.path.join(FRONTEND_DIR, rel_path)

                if os.path.isfile(target_file):
                    self.serve_file(target_file)
                else:
                    self._set_headers(404)
                    self.wfile.write(json.dumps({"error": f"Endpoint or file not found: {path}"}).encode("utf-8"))

        except (ConnectionAbortedError, BrokenPipeError, ConnectionResetError):
            pass
        except Exception as e:
            import traceback
            traceback.print_exc()
            try:
                self._set_headers(500)
                self.wfile.write(json.dumps({"error": str(e)}).encode("utf-8"))
            except Exception:
                pass

    def handle_metadata(self):
        master = data_engine.load_master_data()
        periods = data_engine.get_available_periods()
        gbs = sorted(list({p["gb"] for p in master.values() if p["gb"]}))
        keterangan_opts = data_engine.get_keterangan_options()

        product_options = [
            {
                "code": pcode,
                "name": p["product_name"],
                "gb": p.get("gb", "Unassigned"),
                "keterangan": p.get("keterangan", "Regular"),
                "label": f"{pcode} - {p['product_name']}"
            }
            for pcode, p in sorted(master.items(), key=lambda x: x[0])
        ]

        res = {
            "periods": periods,
            "gb_options": gbs,
            "keterangan_options": keterangan_opts,
            "product_options": product_options,
            "avg_months_options": [1, 3, 6, 9, 12, 2025, 2026],
            "total_products": len(master)
        }
        self._set_headers(200)
        self.wfile.write(json.dumps(res).encode("utf-8"))

    def get_filtered_data(self, get_param):
        period = get_param("period", "")
        avg_months = int(get_param("avg_months", "6"))
        view_mode = get_param("view_mode", get_param("view", get_param("tab", "combined"))).lower()
        
        report = data_engine.get_doi_mnj_report(period=period if period else None, avg_months=avg_months)

        gb_raw = get_param("gb", "All")
        gb_set = parse_multi_param(gb_raw)

        ket_raw = get_param("keterangan", "All")
        ket_set = parse_multi_param(ket_raw)

        prod_raw = get_param("products", get_param("product", "All"))
        prod_set = parse_multi_param(prod_raw)

        health_status = get_param("health_status", "All")
        search = get_param("search", "").strip().lower()

        filtered = []
        for r in report:
            if gb_set and r["gb"] not in gb_set:
                continue

            if ket_set and r["keterangan_produk"] not in ket_set:
                continue

            p_code = r.get("product_code", "")
            p_pcode = r.get("principal_product_code", "")
            p_old = r.get("old_code", "")
            if prod_set and p_code not in prod_set and p_pcode not in prod_set and p_old not in prod_set:
                continue

            if health_status != "All":
                status_check = r.get("health_status_mnj") if view_mode == "mnj" else (r.get("health_status_kx") if view_mode == "kx" else r.get("health_status_total"))
                if status_check != health_status:
                    continue

            if search:
                code_match = search in r["product_code"].lower()
                pcode_match = search in r["principal_product_code"].lower()
                name_match = search in r["product_name"].lower()
                if not (code_match or pcode_match or name_match):
                    continue

            filtered.append(r)

        # Calculate selisih GB: (selisih stok item / total avg sales GB) * 30.0
        # Catatan: Selisih GB & urutan GB selalu menggunakan mode 'value' (Rupiah), meskipun unit='qty' dipilih
        if filtered:
            gb_summary_val = data_engine.get_gb_summary_report(
                period=period if period else None,
                avg_months=avg_months,
                keterangan=ket_raw,
                unit="value",
                products=prod_raw,
                health_status=health_status,
                view_mode=view_mode
            )
            
            is_gb_active = bool(gb_set and "all" not in [g.lower() for g in gb_set])

            for r in filtered:
                ket_p = (r.get("keterangan_produk") or "").strip().lower()
                if view_mode == "mnj":
                    selisih_stok_val = r.get("stok_mnj_value", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_value_mnj", 0.0)
                    selisih_stok_qty = r.get("stok_mnj_qty", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_qty_mnj", 0.0)
                elif view_mode == "kx":
                    selisih_stok_val = r.get("stok_kx_value", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_value_kx", 0.0)
                    selisih_stok_qty = r.get("stok_kx_qty", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_qty_kx", 0.0)
                else:
                    selisih_stok_val = r.get("stok_total_value", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_value", 0.0)
                    selisih_stok_qty = r.get("stok_total_qty", 0.0) if ket_p in ["streamline", "festive"] else r.get("selisih_qty", 0.0)

                if is_gb_active:
                    gb_sales_map = {g["gb"]: g["avg_sales_value"] for g in gb_summary_val}
                    r_gb = r.get("gb", "Unassigned")
                    gb_avg_sales = gb_sales_map.get(r_gb, 0.0)
                    selisih_gb_days = round((selisih_stok_val / gb_avg_sales * 30.0), 2) if gb_avg_sales > 0 else 0.0
                else:
                    total_avg_sales_all_gb = sum(g["avg_sales_value"] for g in gb_summary_val)
                    selisih_gb_days = round((selisih_stok_val / total_avg_sales_all_gb * 30.0), 2) if total_avg_sales_all_gb > 0 else 0.0

                r["selisih_doi_gb"] = selisih_gb_days
                r["selisih_value_gb"] = selisih_stok_val
                r["selisih_qty_gb"] = selisih_stok_qty

        # Selalu urutkan data berdasarkan Selisih Stok (Value/Rupiah) terbesar ke kecil (descending), baik filter GB All maupun spesifik GB
        if filtered:
            def get_sort_key(x):
                ket_p = (x.get("keterangan_produk") or "").strip().lower()
                if view_mode == "mnj":
                    return x.get("stok_mnj_value", 0.0) if ket_p in ["streamline", "festive"] else x.get("selisih_value_mnj", 0.0)
                elif view_mode == "kx":
                    return x.get("stok_kx_value", 0.0) if ket_p in ["streamline", "festive"] else x.get("selisih_value_kx", 0.0)
                else:
                    return x.get("stok_total_value", 0.0) if ket_p in ["streamline", "festive"] else x.get("selisih_value", 0.0)

            filtered.sort(key=get_sort_key, reverse=True)

        return filtered

    def handle_summary(self, get_param):
        filtered = self.get_filtered_data(get_param)
        unit = get_param("unit", "qty").lower()
        view_mode = get_param("view_mode", get_param("view", get_param("tab", "combined"))).lower()

        under = 0
        normal = 0
        over = 0
        
        tot_mnj_val = 0.0
        tot_mnj_qty = 0.0
        tot_kx_val = 0.0
        tot_kx_qty = 0.0
        tot_comb_val = 0.0
        tot_comb_qty = 0.0
        tot_sales_val = 0.0
        tot_sales_qty = 0.0

        for r in filtered:
            status = r.get("health_status_mnj") if view_mode == "mnj" else (r.get("health_status_kx") if view_mode == "kx" else r.get("health_status_total"))
            if status == "Understock":
                under += 1
            elif status == "Normal":
                normal += 1
            elif status == "Overstock":
                over += 1

            tot_mnj_val += r["stok_mnj_value"]
            tot_mnj_qty += r["stok_mnj_qty"]

            tot_kx_val += r["stok_kx_value"]
            tot_kx_qty += r["stok_kx_qty"]

            tot_comb_val += r["stok_total_value"]
            tot_comb_qty += r["stok_total_qty"]

            tot_sales_val += r["avg_sales_value"]
            tot_sales_qty += r["avg_sales_qty"]

        period_active = filtered[0]["period"] if filtered else get_param("period", "2026-07")

        res = {
            "period": period_active,
            "unit": unit,
            "total_sku": len(filtered),
            "understock_count": under,
            "normal_count": normal,
            "overstock_count": over,

            # MNJ Stock
            "total_stok_mnj_value": round(tot_mnj_val, 2),
            "total_stok_mnj_qty": round(tot_mnj_qty, 2),

            # KX Stock
            "total_stok_kx_value": round(tot_kx_val, 2),
            "total_stok_kx_qty": round(tot_kx_qty, 2),

            # Combined Stock
            "total_stok_combined_value": round(tot_comb_val, 2),
            "total_stok_combined_qty": round(tot_comb_qty, 2),

            # Sales
            "total_avg_sales_value": round(tot_sales_val, 2),
            "total_avg_sales_qty": round(tot_sales_qty, 2)
        }
        self._set_headers(200)
        self.wfile.write(json.dumps(res).encode("utf-8"))

    def handle_gb_summary(self, get_param):
        period = get_param("period", "")
        avg_months = int(get_param("avg_months", "6"))
        keterangan = get_param("keterangan", "All")
        products = get_param("products", "All")
        health_status = get_param("health_status", "All")
        unit = get_param("unit", "value")
        view_mode = get_param("view_mode", get_param("view", get_param("tab", "combined"))).lower()
        gb_raw = get_param("gb", "All")
        gb_set = parse_multi_param(gb_raw)

        gb_summary = data_engine.get_gb_summary_report(
            period=period if period else None,
            avg_months=avg_months,
            keterangan=keterangan,
            unit=unit,
            products=products,
            health_status=health_status,
            view_mode=view_mode
        )

        if gb_set and "all" not in [g.lower() for g in gb_set]:
            gb_summary = [g for g in gb_summary if g["gb"] in gb_set]

        self._set_headers(200)
        self.wfile.write(json.dumps(gb_summary).encode("utf-8"))

    def handle_category_gb_summary(self, get_param):
        period = get_param("period", "")
        avg_months = int(get_param("avg_months", "6"))
        keterangan = get_param("keterangan", "All")
        products = get_param("products", "All")
        health_status = get_param("health_status", "All")
        unit = get_param("unit", "value")
        view_mode = get_param("view_mode", get_param("view", get_param("tab", "combined"))).lower()
        gb_raw = get_param("gb", "All")
        gb_set = parse_multi_param(gb_raw)

        cat_summary = data_engine.get_category_gb_summary_report(
            period=period if period else None,
            avg_months=avg_months,
            keterangan=keterangan,
            unit=unit,
            products=products,
            health_status=health_status,
            gb=gb_raw,
            view_mode=view_mode
        )

        if gb_set and "all" not in [g.lower() for g in gb_set]:
            cat_summary = [g for g in cat_summary if g["gb"] in gb_set]

        self._set_headers(200)
        self.wfile.write(json.dumps(cat_summary).encode("utf-8"))

    def handle_doi_trend(self, get_param):
        gb = get_param("gb", "All")
        keterangan = get_param("keterangan", "All")
        products = get_param("products", "All")
        health_status = get_param("health_status", "All")
        avg_months = int(get_param("avg_months", "6"))
        unit = get_param("unit", "value")
        period = get_param("period", get_param("until_period", ""))
        view_mode = get_param("view_mode", get_param("view", get_param("tab", "combined"))).lower()

        trend_data = data_engine.get_historical_doi_trend(
            gb=gb,
            keterangan=keterangan,
            avg_months=avg_months,
            unit=unit,
            products=products,
            health_status=health_status,
            until_period=period if period else None,
            view_mode=view_mode
        )

        self._set_headers(200)
        self.wfile.write(json.dumps(trend_data).encode("utf-8"))

    def handle_doi_data(self, get_param):
        filtered = self.get_filtered_data(get_param)
        page = int(get_param("page", "1"))
        page_size = int(get_param("page_size", "50"))

        total_records = len(filtered)
        total_pages = max(1, (total_records + page_size - 1) // page_size)
        
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        page_data = filtered[start_idx:end_idx]

        res = {
            "total_records": total_records,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages,
            "data": page_data
        }
        self._set_headers(200)
        self.wfile.write(json.dumps(res).encode("utf-8"))

    def handle_export(self, get_param):
        filtered = self.get_filtered_data(get_param)
        
        output = io.StringIO()
        writer = csv.writer(output)
        
        writer.writerow([
            "Periode", "Product Code", "Principal Product Code", "Product Name", "GB", "Keterangan Produk",
            "Harga Dasar (IDR)", "Qty Baik MNJ", "Qty BDP MNJ", "Stok MNJ Qty", "Stok MNJ Value (IDR)", "DOI MNJ (Hari)",
            "Stok KX Qty", "Stok KX Value (IDR)", "DOI KX (Hari)",
            "Total Combined Stok Qty", "Total Combined Stok Value (IDR)", "DOI Total (Hari)", "DOI Min (Hari)", "DOI Max (Hari)",
            "Selisih DOI (Hari)", "Selisih Valuasi (IDR)", "DOI setelah Selisih (Hari)", "Kelebihan Overstock (IDR)", "Defisit Understock (IDR)",
            "Avg Sales Qty", "Avg Sales Value (IDR)", "Status Health Total"
        ])

        for r in filtered:
            writer.writerow([
                r["period"], r["product_code"], r["principal_product_code"], r["product_name"],
                r["gb"], r["keterangan_produk"], r["harga_dasar"],
                r["qty_baik"], r["qty_bdp"], r["stok_mnj_qty"], r["stok_mnj_value"], r["doi_mnj_days"],
                r["stok_kx_qty"], r["stok_kx_value"], r["doi_kx_days"],
                r["stok_total_qty"], r["stok_total_value"], r["doi_total_days"], r["doi_min_days"], r["doi_max_days"],
                r.get("selisih_doi_days", 0.0), r.get("selisih_value", 0.0), r.get("doi_after_selisih", 0.0), r.get("value_overstock", 0.0), r.get("value_understock", 0.0),
                r["avg_sales_qty"], r["avg_sales_value"], r["health_status_total"]
            ])

        csv_content = output.getvalue()
        
        self.send_response(200)
        self.send_header("Content-Type", "text/csv")
        self.send_header("Content-Disposition", "attachment; filename=doi_mnj_kx_report.csv")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(csv_content.encode("utf-8"))

    def handle_export_excel(self, get_param):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        filtered = self.get_filtered_data(get_param)
        period_str = get_param("period", "2026-07")

        wb = openpyxl.Workbook()

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def style_header(ws, headers):
            ws.append(headers)
            ws.row_dimensions[1].height = 28
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

        def auto_fit_columns(ws):
            ws.views.sheetView[0].showGridLines = True
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if '\n' in val_str:
                        val_str = max(val_str.split('\n'), key=len)
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Sheet 1: Total Konsolidasi
        ws_comb = wb.active
        ws_comb.title = "Total Konsolidasi"
        headers_comb = [
            "Periode", "Kode Produk", "Kode Principal", "Nama Produk", "GB", "Keterangan",
            "Harga Dasar (IDR)", "Stok MNJ Qty", "Stok MNJ Value (IDR)", "DOI MNJ (Hari)",
            "Stok KX Qty", "Stok KX Value (IDR)", "DOI KX (Hari)",
            "Combined Stok Qty", "Combined Stok Value (IDR)", "DOI Total (Hari)",
            "DOI Min (Hari)", "DOI Max (Hari)", "Selisih DOI (Hari)", "Selisih Valuasi (IDR)",
            "Avg Sales Qty", "Avg Sales Value (IDR)", "Status Health Total"
        ]
        style_header(ws_comb, headers_comb)

        for r in filtered:
            ws_comb.append([
                r["period"], r["product_code"], r["principal_product_code"], r["product_name"],
                r["gb"], r["keterangan_produk"], r["harga_dasar"],
                r["stok_mnj_qty"], r["stok_mnj_value"], r["doi_mnj_days"],
                r["stok_kx_qty"], r["stok_kx_value"], r["doi_kx_days"],
                r["stok_total_qty"], r["stok_total_value"], r["doi_total_days"],
                r["doi_min_days"], r["doi_max_days"], r.get("selisih_doi_days", 0.0), r.get("selisih_value", 0.0),
                r["avg_sales_qty"], r["avg_sales_value"], r["health_status_total"]
            ])
        auto_fit_columns(ws_comb)

        # Sheet 2: DOI MNJ (Distributor)
        ws_mnj = wb.create_sheet(title="DOI MNJ (Distributor)")
        headers_mnj = [
            "Periode", "Kode Produk", "Kode Principal", "Nama Produk", "GB", "Keterangan",
            "Harga Dasar (IDR)", "Qty Baik MNJ", "Qty BDP MNJ", "Total Stok MNJ Qty", "Stok MNJ Value (IDR)",
            "Avg Sales Qty", "Avg Sales Value (IDR)", "DOI MNJ (Hari)", "DOI Max Master (Hari)",
            "Selisih DOI MNJ (Hari)", "Selisih Stok MNJ (IDR)", "Status Stok MNJ"
        ]
        style_header(ws_mnj, headers_mnj)

        for r in filtered:
            ket_p = (r.get("keterangan_produk") or "").strip().lower()
            is_spec = ket_p in ["streamline", "festive"]
            doi_max = r["doi_max_days"]
            doi_mnj = r["doi_mnj_days"]
            sel_doi_mnj = round(max(0.0, doi_mnj - doi_max), 1) if (not is_spec and doi_max > 0 and doi_mnj > doi_max) else (
                round(doi_mnj - doi_max, 1) if (not is_spec and doi_max > 0 and doi_mnj < r["doi_min_days"]) else 0.0
            )
            sel_stok_mnj = r["stok_mnj_value"] if is_spec else (
                round(max(0.0, r["stok_mnj_value"] - (doi_max / 30.0 * r["avg_sales_value"])), 2) if (doi_max > 0 and doi_mnj > doi_max) else 0.0
            )

            ws_mnj.append([
                r["period"], r["product_code"], r["principal_product_code"], r["product_name"],
                r["gb"], r["keterangan_produk"], r["harga_dasar"],
                r["qty_baik"], r["qty_bdp"], r["stok_mnj_qty"], r["stok_mnj_value"],
                r["avg_sales_qty"], r["avg_sales_value"], doi_mnj, doi_max,
                sel_doi_mnj, sel_stok_mnj, r["health_status_mnj"]
            ])
        auto_fit_columns(ws_mnj)

        # Sheet 3: DOI KX (Principal)
        ws_kx = wb.create_sheet(title="DOI KX (Principal)")
        headers_kx = [
            "Periode", "Kode Produk", "Kode Principal", "Nama Produk", "GB", "Keterangan",
            "Harga Dasar (IDR)", "Stok KX Qty", "Stok KX Value (IDR)",
            "Avg Sales Qty", "Avg Sales Value (IDR)", "DOI KX (Hari)", "DOI Max Master (Hari)",
            "Selisih DOI KX (Hari)", "Selisih Stok KX (IDR)", "Status Stok KX"
        ]
        style_header(ws_kx, headers_kx)

        for r in filtered:
            ket_p = (r.get("keterangan_produk") or "").strip().lower()
            is_spec = ket_p in ["streamline", "festive"]
            doi_max = r["doi_max_days"]
            doi_kx = r["doi_kx_days"]
            sel_doi_kx = round(max(0.0, doi_kx - doi_max), 1) if (not is_spec and doi_max > 0 and doi_kx > doi_max) else (
                round(doi_kx - doi_max, 1) if (not is_spec and doi_max > 0 and doi_kx < r["doi_min_days"]) else 0.0
            )
            sel_stok_kx = r["stok_kx_value"] if is_spec else (
                round(max(0.0, r["stok_kx_value"] - (doi_max / 30.0 * r["avg_sales_value"])), 2) if (doi_max > 0 and doi_kx > doi_max) else 0.0
            )

            ws_kx.append([
                r["period"], r["product_code"], r["principal_product_code"], r["product_name"],
                r["gb"], r["keterangan_produk"], r["harga_dasar"],
                r["stok_kx_qty"], r["stok_kx_value"],
                r["avg_sales_qty"], r["avg_sales_value"], doi_kx, doi_max,
                sel_doi_kx, sel_stok_kx, r["health_status_kx"]
            ])
        auto_fit_columns(ws_kx)

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_bytes = excel_buffer.getvalue()

        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f"attachment; filename=Laporan_MultiSheet_DOI_MNJ_KX_{period_str}.xlsx")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(excel_bytes)

    def handle_export_ppt(self, get_param):
        filters = {
            "period": get_param("period", "2026-07"),
            "unit": get_param("unit", "value"),
            "gb": get_param("gb", "All"),
            "keterangan": get_param("keterangan", "All"),
            "products": get_param("products", "All"),
            "health_status": get_param("health_status", "All"),
            "avg_months": int(get_param("avg_months", "6"))
        }

        template_file = os.path.join(BASE_DIR, "Template PPT.pptx")
        ppt_buffer = generate_doi_ppt(data_engine, filters, template_path=template_file)
        ppt_bytes = ppt_buffer.getvalue()

        period_str = filters["period"]
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.send_header("Content-Disposition", f"attachment; filename=Laporan_DOI_MNJ_KX_{period_str}.pptx")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(ppt_bytes)

def run_server(port=8000):
    print("[SERVER] Preloading datasets into memory...", flush=True)
    data_engine.preload_all_data()

    server_address = ("0.0.0.0", port)
    httpd = ThreadingHTTPServer(server_address, DOIRequestHandler)
    print(f"[SERVER] DOI MNJ & KX Monitoring API Server running on http://localhost:{port} and http://127.0.0.1:{port}", flush=True)

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping server...", flush=True)
        httpd.server_close()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    run_server(port=port)
